"""docx 公式图提取器的单测。

覆盖：
1. 合成 docx（文字 + 内嵌图）能抽出图并留下占位符；
2. 表格里「只含图片」的单元格与旧管线行为一致（丢弃）；
3. 去占位符后的文本与旧管线 ``p.text.strip()`` 逐字一致；
4. plain ↔ marked 双向位置映射自洽。
"""
from __future__ import annotations

import os
import zipfile

import pytest

from services.docx_formula import (
    MARKER_RE,
    AlignedText,
    build_aligned,
    formula_marker,
    marker_positions,
    parse_docx_with_images,
    strip_markers,
)

docx = pytest.importorskip("docx")


def _make_png(path: str) -> str:
    from PIL import Image

    Image.new("RGB", (40, 20), (10, 20, 30)).save(path)
    return path


def _build_sample_docx(tmp_path, with_table: bool = False) -> str:
    """造一个含「文字 + 内嵌公式图」的 docx。"""
    import docx as _docx
    from docx.shared import Inches

    png = _make_png(str(tmp_path / "eq.png"))
    d = _docx.Document()
    d.add_paragraph("第一段普通文字")
    p = d.add_paragraph()
    p.add_run("设函数")
    p.add_run().add_picture(png, width=Inches(0.4))
    p.add_run("连续，则")
    d.add_paragraph("第二段普通文字")
    if with_table:
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "表头A"
        t.cell(0, 1).text = "表头B"
        t.cell(1, 0).text = "值1"
        t.cell(1, 1).text = "值2"
    out = str(tmp_path / "sample.docx")
    d.save(out)
    return out


def test_parse_extracts_image_and_markers(tmp_path):
    path = _build_sample_docx(tmp_path)
    res = parse_docx_with_images(path, out_dir=str(tmp_path / "out"))

    assert res.error == ""
    assert res.paragraph_count >= 3
    assert len(res.images) == 1
    img = res.images[0]
    assert img.index == 0
    assert img.byte_size > 0
    assert img.out_path and os.path.exists(img.out_path)
    assert img.width_px == 40 and img.height_px == 20

    assert "第一段普通文字" in res.text
    assert formula_marker(0) in res.text
    # 占位符必须落在文字中间，不能跑到段首或段尾
    line = [x for x in res.text.split("\n") if formula_marker(0) in x][0]
    assert line.startswith("设函数")
    assert line.endswith("连续，则")


def test_marker_index_is_document_order(tmp_path):
    path = _build_sample_docx(tmp_path)
    res = parse_docx_with_images(path)
    assert res.markers == [0]
    assert res.image_by_index(0) is not None
    assert res.image_by_index(99) is None


def test_strip_markers_and_positions():
    text = "设a[[F:3]]则b"
    assert strip_markers(text) == "设a则b"
    assert marker_positions(text) == [(2, 9)]
    assert MARKER_RE.findall(text) == ["3"]


def test_aligned_plain_matches_legacy_paragraph_text(tmp_path):
    """去占位符后的文本必须与旧管线（p.text.strip()）逐字一致。"""
    import docx as _docx

    path = _build_sample_docx(tmp_path)
    res = parse_docx_with_images(path)
    aligned = build_aligned(res)

    legacy = _docx.Document(path)
    legacy_parts = [p.text.strip() for p in legacy.paragraphs if p.text.strip()]
    assert aligned.plain == "\n".join(legacy_parts)


def test_aligned_mapping_roundtrip(tmp_path):
    path = _build_sample_docx(tmp_path)
    res = parse_docx_with_images(path)
    aligned = build_aligned(res)

    marked = aligned.marked
    # marked 的每个占位符，反向映射到 plain 后应落在占位符被抹掉的位置
    for m in MARKER_RE.finditer(marked):
        p1 = aligned.marked_to_plain(m.start())
        p2 = aligned.marked_to_plain(m.end())
        assert p1 <= p2
        assert aligned.plain[p1:p2] == ""

    # plain 的某个位置正向映射回 marked，应落在同一行内
    probe = aligned.plain.find("连续")
    if probe >= 0:
        mpos = aligned.plain_to_marked(probe)
        assert marked[mpos:mpos + 2] == "连续"


def test_table_cells_without_text_are_dropped(tmp_path):
    """只含公式图的表格单元格要像旧管线那样被丢弃，保证对齐。"""
    import docx as _docx
    from docx.shared import Inches

    png = _make_png(str(tmp_path / "eq2.png"))
    d = _docx.Document()
    d.add_paragraph("正文")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "有文字"
    cell = t.cell(0, 1)
    cell.paragraphs[0].add_run().add_picture(png, width=Inches(0.3))
    out = str(tmp_path / "tbl.docx")
    d.save(out)

    res = parse_docx_with_images(out)
    aligned = build_aligned(res)
    legacy = _docx.Document(out)
    legacy_parts = [p.text.strip() for p in legacy.paragraphs if p.text.strip()]
    for table in legacy.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                legacy_parts.append(" | ".join(cells))
    assert aligned.plain == "\n".join(legacy_parts)


def test_missing_file_returns_error_not_crash(tmp_path):
    res = parse_docx_with_images(str(tmp_path / "不存在.docx"))
    assert res.error
    assert res.text == ""


def test_empty_result_alignment_is_safe():
    from services.docx_formula import DocxParseResult

    aligned = build_aligned(DocxParseResult())
    assert aligned.plain == ""
    assert aligned.marked == ""
    assert aligned.plain_to_marked(0) == 0
    assert aligned.marked_to_plain(0) == 0
